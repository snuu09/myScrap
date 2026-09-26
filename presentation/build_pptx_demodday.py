#!/usr/bin/env python3
"""Build the Demo Day MyBrary deck. Does not overwrite mybrary-4min.pptx."""

import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from build_pptx import (
    DEMO_URL,
    DARK,
    ICONS,
    INK,
    LINE,
    MAGNET,
    MANILA,
    MUTED,
    PAPER,
    ROOT,
    SHOTS,
    SOFT,
    WHITE,
    blank,
    card,
    chips,
    contain,
    demo_footer,
    kicker,
    link_shape,
    notes,
    paint,
    textbox,
    title,
)

OUT = ROOT / "mybrary-demodday.pptx"
TECH = ICONS / "tech"
DEMO_SHOTS = SHOTS / "demodday"


def bullet_list(slide, items, x, y, w, line_h=0.32, size=13):
    for i, item in enumerate(items):
        textbox(slide, item, x, y + i * line_h, w, line_h, size, True, INK)


def demo_image_card(slide, filename, x, y, w, h, label="", href=None):
    box = card(slide, x, y, w, h)
    contain(slide, DEMO_SHOTS / filename, x + 0.08, y + 0.08, w - 0.16, h - 0.16)
    if label:
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + w - 1.5), Inches(y + 0.12), Inches(1.25), Inches(0.36)
        )
        pill.adjustments[0] = 0.45
        paint(pill, MAGNET)
        textbox(slide, label, x + w - 1.42, y + 0.2, 1.1, 0.18, 8, True, WHITE, "center")
    if href:
        link_shape(box, href)
    return box


def tech_card(slide, icon_stem, label, sub, x, y, w, h):
    shape = card(slide, x, y, w, h)
    icon = TECH / f"{icon_stem}.png"
    slide.shapes.add_picture(str(icon), Inches(x + 0.22), Inches(y + (h - 0.55) / 2), Inches(0.55), Inches(0.55))
    textbox(slide, label, x + 0.95, y + 0.28, w - 1.15, 0.32, 16, True, INK)
    textbox(slide, sub, x + 0.95, y + 0.62, w - 1.15, 0.28, 12, False, MUTED)
    return shape


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "MyBrary Demo Day"
    prs.core_properties.subject = "붙여넣으면, AI가 분류해 줘요"

    # 1 cover
    s = blank(prs)
    textbox(s, "성남청년 AI 캠프 · Demo Day", 0.7, 1.55, 11, 0.35, 15, True, MAGNET)
    textbox(s, "2026.09.19", 0.7, 1.95, 11, 0.3, 14, True, MUTED)
    textbox(s, "MyBrary", 0.7, 2.35, 11, 1.0, 56, True, INK)
    textbox(s, "붙여넣으면, AI가 분류해 줘요.", 0.7, 3.45, 11, 0.5, 24, False, SOFT)
    textbox(s, "박경진", 0.7, 4.05, 11, 0.4, 20, True, INK)
    chips(
        s,
        (("붙여넣기", "paste", 1.75), ("AI 분류", "sparkle", 1.65), ("다시 찾기", "search", 1.85)),
        0.7,
        4.7,
    )
    textbox(s, "개인 스크랩 · 글 · 사진 · 영상 · 소리 · 링크 · 문서", 0.7, 6.2, 11, 0.4, 15, False, MUTED)
    notes(s, "Demo Day입니다. 붙여넣으면 AI가 분류해 주는 개인 책장입니다.")

    # 2 problem: tab narrative + existing scrap friction
    s = blank(prs)
    kicker(s, "문제 · 아이디어")
    title(s, "탭은 쌓이는데, 다시 보기 어렵다.", y=0.52, size=26)
    textbox(
        s,
        "브라우저는 탭으로, 스크랩 서비스는 형태별로 갈라집니다. 모아 두고도 다시 찾기 어렵습니다.",
        0.55,
        1.12,
        12.2,
        0.35,
        13,
        False,
        MUTED,
    )
    visual = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.55), Inches(6.4), Inches(5.4))
    visual.adjustments[0] = 0.04
    paint(visual, DARK)
    contain(s, SHOTS / "desktop-intro.png", 0.7, 1.7, 6.1, 5.1)

    icon_rows = (
        (
            "pages",
            "누구",
            "탭을 많이 쓰고, 스크랩도 자주 하는 사람",
            "글·링크·파일·영상을 열어두고 저장하지만, 도구가 여러 개로 갈라집니다.",
        ),
        (
            "clipboard",
            "불편 · 탭",
            "쌓인 탭은 다음에 확인하기 어렵다",
            "카테고리도 없어 더 안 보게 되고, 다시 열어볼 흐름이 없습니다.",
        ),
        (
            "file",
            "불편 · 스크랩",
            "서비스는 있어도 형태·확인이 불편함",
            "일부 자료만 받거나, 링크는 직접 열어봐야 내용을 확인할 수 있습니다.",
        ),
        (
            "search",
            "왜 안 풀리나",
            "저장과 확인이 따로 움직임",
            "형태마다 앱이 갈라지고, 미리보기·요약이 약해 원문으로 다시 이동해야 합니다.",
        ),
        (
            "sparkle",
            "그래서 MyBrary",
            "붙여넣으면 초안이 나와요",
            "여러 형태를 한 책장에 모으고, AI 분류·요약·태그로 다시 찾아 미리보기로 이동 없이 확인합니다.",
        ),
    )
    for i, (icon, label, head, body) in enumerate(icon_rows):
        y = 1.55 + i * 1.08
        fill = MANILA if i == 4 else PAPER
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(y), Inches(5.6), Inches(1.0))
        shape.adjustments[0] = 0.06
        paint(shape, fill, None if i == 4 else LINE)
        s.shapes.add_picture(str(ICONS / (icon + ".png")), Inches(7.4), Inches(y + 0.14), Inches(0.24), Inches(0.24))
        textbox(s, label, 7.75, y + 0.14, 4.8, 0.22, 10, True, MAGNET)
        textbox(s, head, 7.4, y + 0.38, 5.2, 0.24, 12, True, INK)
        textbox(s, body, 7.4, y + 0.62, 5.2, 0.32, 10, False, MUTED)
    notes(
        s,
        "탭이 쌓이면 다시 보기 어렵고 카테고리도 없습니다. 스크랩 서비스도 형태·확인이 불편하고, "
        "저장과 확인이 따로 움직입니다. MyBrary는 붙여넣으면 초안이 나오고 AI 분류·요약·태그로 다시 찾게 합니다.",
    )

    # 3 AI drafts
    s = blank(prs)
    kicker(s, "AI 자동 분류")
    title(s, "붙여넣으면, 결과 초안이 먼저 나와요.", y=0.52, size=26)
    chips(
        s,
        (("제목", "text", 1.15), ("태그", "tag", 1.15), ("분류", "bookmark", 1.2), ("요약·분석", "sparkle", 1.7)),
        0.55,
        1.2,
    )
    drafts = (
        ("desktop-draft-memo.png", "메모 초안"),
        ("desktop-draft-link.png", "링크 초안"),
        ("desktop-draft-video.png", "영상 초안"),
    )
    for i, (filename, label) in enumerate(drafts):
        x = 0.5 + i * 3.35
        demo_image_card(s, filename, x, 1.85, 3.2, 4.55, label)
    phone_x = 10.7
    frame = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(phone_x), Inches(1.75), Inches(2.2), Inches(4.75))
    frame.adjustments[0] = 0.08
    paint(frame, DARK)
    contain(s, SHOTS / "mobile-plus.png", phone_x + 0.1, 1.9, 2.0, 4.45)
    demo_footer(s)
    notes(s, "웹 기준으로 메모·링크·영상 초안을 보여 주세요. 폰은 부가적으로, 같은 흐름이 모바일에서도 됩니다.")

    # 4 supported types
    s = blank(prs)
    kicker(s, "핵심 기능 · 지원 형태")
    title(s, "여러 형태를, 같은 책장에 모아요.", y=0.52, size=26)
    chips(
        s,
        (
            ("PDF·문서", "file", 1.55),
            ("링크", "link", 1.2),
            ("영상", "video", 1.2),
            ("소리", "audio", 1.2),
            ("메모", "text", 1.15),
            ("사진", "image", 1.15),
        ),
        0.55,
        1.2,
    )
    type_shots = (
        ("desktop-detail-pdf.png", "PDF"),
        ("desktop-detail-link.png", "링크"),
        ("desktop-detail-video.png", "영상"),
        ("desktop-detail-audio.png", "소리"),
    )
    for i, (filename, label) in enumerate(type_shots):
        x = 0.5 + (i % 2) * 5.0
        y = 1.85 + (i // 2) * 2.45
        demo_image_card(s, filename, x, y, 4.85, 2.3, label, href=DEMO_URL)
    phone_x = 10.7
    frame = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(phone_x), Inches(1.75), Inches(2.2), Inches(4.95))
    frame.adjustments[0] = 0.08
    paint(frame, DARK)
    contain(s, DEMO_SHOTS / "mobile-detail.png", phone_x + 0.1, 1.9, 2.0, 4.65)
    demo_footer(s)
    notes(
        s,
        "지원 형태를 보여 준 뒤, 하단 주소나 스크린샷을 눌러 라이브 웹을 엽니다. https://mybrary-snuu09.web.app/",
    )

    # 5 MVP result — equal-level bullet lists
    s = blank(prs)
    kicker(s, "MVP 결과")
    title(s, "핵심 흐름은 돌아가고, 관리까지 붙었어요.")
    demo_image_card(s, "desktop-mvp-gallery.png", 0.55, 1.45, 4.55, 5.25, "책장", href=DEMO_URL)
    demo_image_card(s, "desktop-mvp-bundled.png", 5.3, 1.45, 3.7, 2.45, "묶음 보기", href=DEMO_URL)
    demo_image_card(s, "desktop-mvp-dashboard.png", 5.3, 4.1, 3.7, 2.6, "대시보드", href=DEMO_URL)

    done_items = (
        "붙여넣기",
        "AI 분류",
        "AI 요약",
        "태깅",
        "검색",
        "상세",
        "책장",
        "대시보드",
        "로그인",
    )
    soon_items = (
        "개인 스토리지(Google Drive) 연동",
        "등급별 플랜 기능 제한",
        "외부 공유",
        "커뮤니티",
    )
    card(s, 9.25, 1.45, 3.55, 2.55)
    textbox(s, "구현됨", 9.45, 1.58, 3.15, 0.28, 12, True, MAGNET)
    bullet_list(s, done_items, 9.45, 1.92, 3.15, line_h=0.22, size=12)

    card(s, 9.25, 4.15, 3.55, 2.55)
    textbox(s, "보완 예정", 9.45, 4.28, 3.15, 0.28, 12, True, MAGNET)
    bullet_list(s, soon_items, 9.45, 4.7, 3.15, line_h=0.36, size=12)

    demo_footer(s)
    notes(s, "MVP 화면을 보여 준 뒤, 하단 주소나 스크린샷을 눌러 라이브 웹을 엽니다. https://mybrary-snuu09.web.app/")

    # 6 tech stack
    s = blank(prs)
    kicker(s, "기술 스택")
    title(s, "언어 · 프레임워크 · 솔루션")
    textbox(s, "언어", 0.7, 1.4, 12, 0.3, 14, True, MAGNET)
    tech_card(s, "typescript", "TypeScript", "언어", 0.7, 1.8, 4.0, 1.15)

    textbox(s, "프레임워크", 0.7, 3.2, 12, 0.3, 14, True, MAGNET)
    tech_card(s, "react", "React", "UI 프레임워크", 0.7, 3.6, 3.9, 1.15)
    tech_card(s, "vite", "Vite", "빌드 도구", 4.8, 3.6, 3.9, 1.15)
    tech_card(s, "tailwindcss", "Tailwind CSS", "스타일", 8.9, 3.6, 3.7, 1.15)

    textbox(s, "솔루션", 0.7, 5.0, 12, 0.3, 14, True, MAGNET)
    tech_card(s, "supabase", "Supabase", "Auth · DB · Storage · Edge", 0.7, 5.4, 4.0, 1.25)
    tech_card(s, "firebase", "Firebase Hosting", "배포", 4.9, 5.4, 4.0, 1.25)
    tech_card(s, "anthropic", "Claude API", "AI 분류 · 요약", 9.1, 5.4, 3.5, 1.25)
    notes(s, "TypeScript, React, Vite, Tailwind, Supabase, Firebase Hosting, Claude API를 사용합니다.")

    # 7 close
    s = blank(prs)
    textbox(s, "감사합니다", 0.7, 2.05, 11, 0.35, 14, True, MAGNET)
    textbox(s, "한번 붙여넣어 보세요.", 0.7, 2.55, 7.6, 1.0, 42, True, INK)
    url_box = textbox(s, "mybrary-snuu09.web.app", 0.7, 3.75, 7.5, 0.55, 24, True, MAGNET)
    link_shape(url_box)
    textbox(s, "QR로 바로 둘러볼 수 있습니다.", 0.7, 4.6, 7.5, 0.4, 16, False, MUTED)
    card(s, 9.35, 1.85, 2.7, 2.7)
    contain(s, SHOTS / "site-qr.png", 9.55, 2.05, 2.3, 2.3)
    notes(s, "라이브에서 둘러보기로 바로 열 수 있습니다.")

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
