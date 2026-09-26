#!/usr/bin/env python3
"""Build the 4-minute MyBrary deck from presentation/shots and icons."""

import struct
from pathlib import Path

from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "shots"
ICONS = ROOT / "icons"
OUT = ROOT / "mybrary-4min.pptx"
FONT = "Apple SD Gothic Neo"

MAGNET = RGBColor(0xE5, 0x6F, 0x0A)
INK = RGBColor(0x32, 0x2C, 0x26)
SOFT = RGBColor(0x52, 0x49, 0x40)
MUTED = RGBColor(0x6E, 0x66, 0x5C)
ENAMEL = RGBColor(0xFF, 0xF6, 0xF0)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE4, 0xD2, 0xC6)
CREAM = RGBColor(0xFF, 0xF1, 0xE4)
MANILA = RGBColor(0xFF, 0xF3, 0xDF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x2A, 0x1A, 0x08)


def png_size(path):
    if path.suffix.lower() == ".png":
        with open(path, "rb") as handle:
            handle.read(16)
            width, height = struct.unpack(">II", handle.read(8))
        return width, height
    with Image.open(path) as img:
        return img.size


def set_font(run, size, bold, color, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", name)


def paint(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)


def textbox(slide, text, x, y, w, h, size, bold, color, align="left"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.paragraphs[0].alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[align]
    run = tf.paragraphs[0].add_run()
    run.text = text
    set_font(run, size, bold, color)
    return box


def notes(slide, line):
    slide.notes_slide.notes_text_frame.text = line


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    paint(bg, ENAMEL)
    sp_tree = slide.shapes._spTree
    sp = bg._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)
    return slide


def chip(slide, label, icon, x, y, w, h=0.46):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.5
    paint(shape, PAPER, LINE)
    icon_size = 0.28
    slide.shapes.add_picture(
        str(ICONS / (icon + ".png")),
        Inches(x + 0.1),
        Inches(y + (h - icon_size) / 2),
        Inches(icon_size),
        Inches(icon_size),
    )
    textbox(slide, label, x + 0.42, y + 0.08, w - 0.52, h - 0.12, 14, True, INK)
    return w


def chips(slide, items, x, y, gap=0.12):
    cursor = x
    for label, icon, width in items:
        chip(slide, label, icon, cursor, y, width)
        cursor += width + gap


def kicker(slide, label):
    textbox(slide, label, 0.55, 0.28, 10, 0.3, 13, True, MAGNET)


def title(slide, label, y=0.58, size=28):
    textbox(slide, label, 0.55, y, 12.2, 0.62, size, True, INK)


def contain(slide, path, x, y, box_w, box_h):
    width, height = png_size(path)
    scale = min(box_w / width, box_h / height)
    draw_w = width * scale
    draw_h = height * scale
    left = x + (box_w - draw_w) / 2
    top = y + (box_h - draw_h) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(draw_w), Inches(draw_h))


def pair(slide, desk, phone, y=1.85, box_h=4.9):
    phone_w = 2.5
    gap = 0.2
    desk_w = 12.3 - phone_w - gap
    desk_card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(y), Inches(desk_w), Inches(box_h)
    )
    desk_card.adjustments[0] = 0.04
    paint(desk_card, PAPER, LINE)
    contain(slide, SHOTS / desk, 0.62, y + 0.12, desk_w - 0.24, box_h - 0.24)
    phone_x = 0.5 + desk_w + gap
    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(phone_x), Inches(y), Inches(phone_w), Inches(box_h)
    )
    frame.adjustments[0] = 0.08
    paint(frame, DARK)
    contain(slide, SHOTS / phone, phone_x + 0.12, y + 0.16, phone_w - 0.24, box_h - 0.32)


def card(slide, x, y, w, h):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.06
    paint(shape, PAPER, LINE)
    return shape


DEMO_URL = "https://mybrary-snuu09.web.app/"
REFLECT_NOTE = (
    "바이브 코딩과 안드로이드 개발만 한 내가 web 에 대한 지식을 경험을 기대하여 지원하였고 "
    "좋은 기회로 인해 2달 간 값진 경험을 하였다. mvp 단계지만 생각한 아이디어를 바이브 코딩을 통해 "
    "구현했고 점점 욕심과 의욕이 생겨 실제 서비스가 될 수 있게 노력할 예정이다"
)


def link_shape(shape, url=DEMO_URL):
    shape.click_action.hyperlink.address = url
    return shape


def demo_footer(slide, y=6.92):
    box = textbox(slide, DEMO_URL, 0.55, y, 12.2, 0.35, 14, True, MAGNET)
    link_shape(box)
    return box


def image_card(slide, filename, x, y, w, h, label="", href=None):
    box = card(slide, x, y, w, h)
    contain(slide, SHOTS / filename, x + 0.08, y + 0.08, w - 0.16, h - 0.16)
    if label:
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + w - 1.5), Inches(y + 0.12), Inches(1.25), Inches(0.36))
        pill.adjustments[0] = 0.45
        paint(pill, MAGNET)
        textbox(slide, label, x + w - 1.42, y + 0.2, 1.1, 0.18, 8, True, WHITE, "center")
    if href:
        link_shape(box, href)
    return box


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "MyBrary 4분"
    prs.core_properties.subject = "붙여넣으면, AI가 분류해 줘요"

    # 1 cover
    s = blank(prs)
    textbox(s, "성남청년 AI 캠프 · MVP", 0.7, 1.75, 11, 0.35, 15, True, MAGNET)
    textbox(s, "MyBrary", 0.7, 2.15, 11, 1.0, 56, True, INK)
    textbox(s, "붙여넣으면, AI가 분류해 줘요.", 0.7, 3.3, 11, 0.5, 24, False, SOFT)
    textbox(s, "박경진", 0.7, 3.9, 11, 0.4, 20, True, INK)
    chips(
        s,
        (("붙여넣기", "paste", 1.75), ("AI 분류", "sparkle", 1.65), ("다시 찾기", "search", 1.85)),
        0.7,
        4.55,
    )
    textbox(s, "개인 스크랩 · 글 · 사진 · 영상 · 소리 · 링크 · 문서", 0.7, 6.2, 11, 0.4, 15, False, MUTED)
    notes(s, "붙여넣으면 AI가 분류해 주는 개인 책장입니다. 4분 이야기입니다.")

    # 2 problem (miricanvas + expanded cards)
    s = blank(prs)
    kicker(s, "문제")
    title(s, "정보는 많은데, 한곳이 없어요.", y=0.52, size=26)
    textbox(
        s,
        "스크랩 서비스는 이미 있습니다. 그래도 모으고 확인하는 흐름은 여전히 불편합니다.",
        0.55,
        1.12,
        12.2,
        0.35,
        13,
        False,
        MUTED,
    )
    visual = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.55), Inches(7.2), Inches(5.4))
    visual.adjustments[0] = 0.04
    paint(visual, DARK)
    contain(s, SHOTS / "problem-miricanvas.png", 0.7, 1.7, 6.9, 5.1)
    icon_rows = (
        ("pages", "누구", "스크랩을 자주 하는 사람", "글·링크·파일·영상을 자주 저장하지만, 쓰는 도구가 여러 개로 갈라집니다."),
        (
            "clipboard",
            "불편",
            "서비스는 있어도 형태·확인이 불편함",
            "스크랩 서비스는 있지만 일부 자료 형태만 받거나, 링크는 직접 열어봐야 내용을 확인할 수 있습니다.",
        ),
        (
            "search",
            "왜 안 풀리나",
            "저장과 확인이 따로 움직임",
            "형태마다 앱이 갈라지고, 미리보기·요약이 약해 원문으로 다시 이동해야 합니다.",
        ),
        (
            "sparkle",
            "그래서 MyBrary",
            "붙여넣으면 초안이 나와요",
            "여러 형태를 한 책장에 모으고, 미리보기·요약으로 이동 없이 확인할 수 있게 합니다.",
        ),
    )
    for i, (icon, label, head, body) in enumerate(icon_rows):
        y = 1.55 + i * 1.35
        fill = MANILA if i == 3 else PAPER
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(y), Inches(4.8), Inches(1.25))
        shape.adjustments[0] = 0.06
        paint(shape, fill, None if i == 3 else LINE)
        s.shapes.add_picture(str(ICONS / (icon + ".png")), Inches(8.2), Inches(y + 0.18), Inches(0.28), Inches(0.28))
        textbox(s, label, 8.6, y + 0.18, 3.9, 0.24, 11, True, MAGNET)
        textbox(s, head, 8.2, y + 0.48, 4.4, 0.28, 14, True, INK)
        textbox(s, body, 8.2, y + 0.78, 4.4, 0.4, 11, False, MUTED)
    notes(
        s,
        "스크랩 서비스는 이미 있습니다. 다만 일부 형태만 받거나, 링크는 직접 열어봐야 합니다. MyBrary는 여러 형태를 한곳에 모으고 미리보기·요약으로 이동 없이 확인합니다.",
    )

    # 3 AI drafts (memo / link / video) + mobile secondary
    s = blank(prs)
    kicker(s, "AI 자동 분류")
    title(s, "붙여넣으면, 결과 초안이 먼저 나와요.", y=0.52, size=26)
    chips(
        s,
        (("제목", "text", 1.15), ("태그", "tag", 1.15), ("분류", "bookmark", 1.2), ("요약·분석", "sparkle", 1.7)),
        0.55,
        1.2,
    )
    drafts = (
        ("desktop-draft-memo.png", "메모 초안"),
        ("desktop-draft-link.png", "링크 초안"),
        ("desktop-draft-video.png", "영상 초안"),
    )
    for i, (filename, label) in enumerate(drafts):
        x = 0.5 + i * 3.35
        image_card(s, filename, x, 1.85, 3.2, 4.55, label)
    phone_x = 10.7
    # 390x844 portrait → ~2.15 x 4.65 in
    frame = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(phone_x), Inches(1.75), Inches(2.2), Inches(4.75))
    frame.adjustments[0] = 0.08
    paint(frame, DARK)
    contain(s, SHOTS / "mobile-plus.png", phone_x + 0.1, 1.9, 2.0, 4.45)
    demo_footer(s)
    notes(s, "웹 기준으로 메모·링크·영상 초안을 보여 주세요. 폰은 부가적으로, 같은 흐름이 모바일에서도 됩니다.")

    # 4 supported types (merged former 4+5)
    s = blank(prs)
    kicker(s, "핵심 기능 · 지원 형태")
    title(s, "여러 형태를, 같은 책장에 모아요.", y=0.52, size=26)
    chips(
        s,
        (
            ("PDF·문서", "file", 1.55),
            ("링크", "link", 1.2),
            ("영상", "video", 1.2),
            ("소리", "audio", 1.2),
            ("메모", "text", 1.15),
            ("사진", "image", 1.15),
        ),
        0.55,
        1.2,
    )
    type_shots = (
        ("desktop-detail-pdf.png", "PDF"),
        ("desktop-detail-link.png", "링크"),
        ("desktop-detail-video.png", "영상"),
        ("desktop-detail-audio.png", "소리"),
    )
    for i, (filename, label) in enumerate(type_shots):
        x = 0.5 + (i % 2) * 5.0
        y = 1.85 + (i // 2) * 2.45
        image_card(s, filename, x, y, 4.85, 2.3, label, href=DEMO_URL)
    phone_x = 10.7
    frame = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(phone_x), Inches(1.75), Inches(2.2), Inches(4.95))
    frame.adjustments[0] = 0.08
    paint(frame, DARK)
    contain(s, SHOTS / "mobile-detail.png", phone_x + 0.1, 1.9, 2.0, 4.65)
    demo_footer(s)
    notes(
        s,
        "지원 형태를 보여 준 뒤, 하단 주소나 스크린샷을 눌러 라이브 웹을 엽니다. https://mybrary-snuu09.web.app/",
    )

    # 5 MVP result
    s = blank(prs)
    kicker(s, "MVP 결과")
    title(s, "핵심 흐름은 돌아가고, 관리까지 붙었어요.")
    image_card(s, "desktop-mvp-gallery.png", 0.55, 1.45, 4.55, 5.25, "책장", href=DEMO_URL)
    image_card(s, "desktop-mvp-bundled.png", 5.3, 1.45, 3.7, 2.45, "묶음 보기", href=DEMO_URL)
    image_card(s, "desktop-mvp-dashboard.png", 5.3, 4.1, 3.7, 2.6, "대시보드", href=DEMO_URL)
    card(s, 9.25, 1.45, 3.55, 2.45)
    textbox(s, "구현됨", 9.5, 1.68, 3, 0.28, 11, True, MAGNET)
    textbox(s, "붙여넣기 · AI 분류 · 책장", 9.5, 2.15, 3.05, 0.5, 15, True, INK)
    textbox(s, "검색 · 상세 · 대시보드 · 로그인 · 배포", 9.5, 2.85, 3.05, 0.6, 12, False, MUTED)
    card(s, 9.25, 4.1, 3.55, 2.6)
    textbox(s, "보완 예정", 9.5, 4.35, 3, 0.28, 11, True, MUTED)
    textbox(s, "결제 · 확장 · 공유", 9.5, 4.85, 3.05, 0.45, 15, True, INK)
    textbox(s, "내보내기 · 네이티브 앱", 9.5, 5.5, 3.05, 0.5, 12, False, MUTED)
    demo_footer(s)
    notes(s, "MVP 화면을 보여 준 뒤, 하단 주소나 스크린샷을 눌러 라이브 웹을 엽니다. https://mybrary-snuu09.web.app/")

    # 6 reflect: camp poster + intro image side by side; copy only in notes
    s = blank(prs)
    kicker(s, "소감")
    left = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.05), Inches(6.0), Inches(5.9))
    left.adjustments[0] = 0.04
    paint(left, DARK)
    contain(s, SHOTS / "camp-poster.jpg", 0.7, 1.2, 5.7, 5.6)
    right = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.05), Inches(6.0), Inches(5.9))
    right.adjustments[0] = 0.04
    paint(right, DARK)
    contain(s, SHOTS / "reflect-intro.jpg", 6.95, 1.2, 5.7, 5.6)
    notes(s, REFLECT_NOTE)

    # 7 close
    s = blank(prs)
    textbox(s, "감사합니다", 0.7, 2.05, 11, 0.35, 14, True, MAGNET)
    textbox(s, "한번 붙여넣어 보세요.", 0.7, 2.55, 7.6, 1.0, 42, True, INK)
    url_box = textbox(s, "mybrary-snuu09.web.app", 0.7, 3.75, 7.5, 0.55, 24, True, MAGNET)
    link_shape(url_box)
    textbox(s, "QR로 바로 둘러볼 수 있습니다.", 0.7, 4.6, 7.5, 0.4, 16, False, MUTED)
    card(s, 9.35, 1.85, 2.7, 2.7)
    contain(s, SHOTS / "site-qr.png", 9.55, 2.05, 2.3, 2.3)
    notes(s, "라이브에서 둘러보기로 바로 열 수 있습니다.")

    prs.save(OUT)
    # keep previous filename as a copy for open tabs
    also = ROOT / "mybrary-5min.pptx"
    also.write_bytes(OUT.read_bytes())
    print(OUT)
    print(also)


if __name__ == "__main__":
    build()
